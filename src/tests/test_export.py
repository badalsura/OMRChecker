"""Export module: profiles, casting, CSV/XLSX/PDF/SQL writers, API and CLI."""

import csv
import io
import json
import random
import sqlite3
from datetime import date
from decimal import Decimal

import cv2
import numpy as np
import pytest

from src.export import ExportError, ExportProfile, LossyCastError, export_results
from src.export.__main__ import main as export_cli
from src.synth.render import default_spec, random_answers, render_sheet
from src.tests.test_api import make_client, png_bytes, wait_for_job

INFO = {
    "output_columns": ["Roll", "q1", "q2", "dob", "agree", "price"],
    "zone_names": ["name"],
    "custom_labels": {"Roll": ["r1", "r2", "r3"]},
}


def fake_result(index, roll="012", q1="A", flagged=False, corrected=False):
    fields = {
        "r1": {"value": roll[0], "confidence": 0.9, "flags": []},
        "r2": {"value": roll[1], "confidence": 0.8, "flags": []},
        "r3": {"value": roll[2], "confidence": 0.95, "flags": []},
        "q1": {
            "value": q1,
            "confidence": 0.4 if flagged else 0.9,
            "flags": ["multi_marked"] if flagged else [],
        },
        "q2": {"value": "B", "confidence": 0.9, "flags": []},
        "dob": {"value": "31122001", "confidence": 0.9, "flags": []},
        "agree": {"value": "Y", "confidence": 0.9, "flags": []},
        "price": {"value": "12.50", "confidence": 0.9, "flags": []},
    }
    if corrected:
        fields["q2"]["original_value"] = "C"
    return {
        "scan_id": f"scan{index:04d}",
        "file_id": f"sheet{index}.png",
        "file_name": f"sheet{index}.png",
        "job_id": "job1",
        "template_id": "t1",
        "page": 0,
        "status": "needs_review" if flagged else "ok",
        "score": 3.5,
        "fields": fields,
        "zones": {"name": {"value": "=HYPERLINK(1)", "confidence": 0.7, "flags": []}},
        "responses": {
            "Roll": roll,
            "q1": q1,
            "q2": "B",
            "dob": "31122001",
            "agree": "Y",
            "price": "12.50",
            "name": "=HYPERLINK(1)",
        },
        "review": [{"name": "q1", "kind": "field", "flags": ["multi_marked"]}]
        if flagged
        else [],
    }


RESULTS = [
    fake_result(0),
    fake_result(1, roll="345", flagged=True),
    fake_result(2, corrected=True),
]


def run(tmp_path, fmt, profile, results=RESULTS, **kwargs):
    out = tmp_path / f"out.{ 'sqlite' if fmt == 'sqlite' else fmt}"
    report = export_results(
        lambda: iter(results),
        fmt,
        out,
        ExportProfile.from_dict(profile),
        [INFO],
        **kwargs,
    )
    return out, report


def test_profile_validation():
    with pytest.raises(ExportError):
        ExportProfile.from_dict({"fields": [{"field": "a", "type": "money"}]})
    with pytest.raises(ExportError):
        ExportProfile.from_dict({"fields": [{"field": "d", "type": "date"}]})
    with pytest.raises(ExportError):
        ExportProfile.from_dict({"meta": ["nope"]})
    with pytest.raises(ExportError):
        ExportProfile.from_dict({"csv": {"leadingZeros": "maybe"}})


def test_csv_types_rename_order_and_leading_zeros(tmp_path):
    profile = {
        "fields": [
            {"field": "q2", "header": "Question 2"},
            {"field": "Roll", "header": "Roll number"},
            {
                "field": "dob",
                "type": "date",
                "format": "%d%m%Y",
                "outputFormat": "%Y-%m-%d",
            },
            {"field": "agree", "type": "bool"},
            {"field": "price", "type": "decimal"},
        ],
        "meta": ["file_name", "scan_id"],
        "includeConfidence": True,
        "includeFlags": True,
        "csv": {"leadingZeros": "formula", "bom": False},
    }
    out, report = run(tmp_path, "csv", profile)
    rows = list(csv.reader(io.StringIO(out.read_text())))
    assert rows[0] == [
        "file_name",
        "scan_id",
        "review_status",
        "corrected",
        "Question 2",
        "Question 2_confidence",
        "Question 2_flags",
        "Roll number",
        "Roll number_confidence",
        "Roll number_flags",
        "dob",
        "dob_confidence",
        "dob_flags",
        "agree",
        "agree_confidence",
        "agree_flags",
        "price",
        "price_confidence",
        "price_flags",
    ]
    first = dict(zip(rows[0], rows[1]))
    assert first["Roll number"] == '="012"'
    assert first["Roll number_confidence"] == "0.8"  # min over the joined parts
    assert first["dob"] == "2001-12-31"
    assert first["agree"] == "true"
    assert first["price"] == "12.50"
    assert first["review_status"] == "auto"
    second = dict(zip(rows[0], rows[2]))
    assert second["Roll number"] == "345"  # no leading zero, left alone
    assert second["review_status"] == "needs_review"
    third = dict(zip(rows[0], rows[3]))
    assert third["corrected"] == "true"
    assert report["rows"] == 3

    profile["csv"]["leadingZeros"] = "apostrophe"
    out, _ = run(tmp_path, "csv", profile)
    assert list(csv.reader(io.StringIO(out.read_text())))[1][7] == "'012"

    # Default: everything, as text, in template order
    out, _ = run(tmp_path, "csv", {"csv": {"bom": False}})
    header = next(csv.reader(io.StringIO(out.read_text())))
    assert header[-7:] == ["Roll", "q1", "q2", "dob", "agree", "price", "name"]


def test_lossy_cast_is_refused_unless_allowed(tmp_path):
    profile = {"fields": [{"field": "Roll", "type": "int"}]}
    with pytest.raises(LossyCastError) as error:
        run(tmp_path, "csv", profile)
    assert "leading zero" in str(error.value) and "allowLossyCast" in str(error.value)
    profile["allowLossyCast"] = True
    out, report = run(tmp_path, "csv", profile)
    rows = list(csv.DictReader(io.StringIO(out.read_text("utf-8-sig"))))
    assert rows[0]["Roll"] == "12"
    assert any("joins 3 fields" in w for w in report["warnings"])
    # Uncastable values (a multi-mark read as int) become empty with a warning
    profile = {"fields": [{"field": "q1", "type": "int"}]}
    out, report = run(tmp_path, "csv", profile)
    assert report["cast_failures"] == 3 and report["warnings"]
    with pytest.raises(LossyCastError):
        run(tmp_path, "csv", {**profile, "strictCast": True})


def test_xlsx_text_cells_highlight_and_sheet_split(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    results = [fake_result(i, flagged=(i == 1), corrected=(i == 2)) for i in range(5)]
    profile = {
        "fields": [
            {"field": "price", "type": "decimal"},
            {"field": "dob", "type": "date", "format": "%d%m%Y"},
        ],
        "includeOtherFields": True,
        "xlsx": {"maxRowsPerSheet": 2},
    }
    out, report = run(tmp_path, "xlsx", profile, results=results)
    assert report["rows"] == 5
    book = openpyxl.load_workbook(out)
    assert book.sheetnames == ["Results", "Results 2", "Results 3"]
    sheet = book["Results"]
    header = [c.value for c in sheet[1]]
    row = {h: c for h, c in zip(header, sheet[2])}
    assert row["Roll"].value == "012" and row["Roll"].data_type == "s"
    assert row["price"].value == 12.5
    assert row["dob"].value.date() == date(2001, 12, 31)
    assert row["name"].value == "=HYPERLINK(1)" and row["name"].data_type == "s"
    flagged = {h: c for h, c in zip(header, sheet[3])}
    assert flagged["q1"].fill.start_color.rgb.endswith("FFF3C2")
    corrected = {h: c for h, c in zip(header, book["Results 2"][2])}
    assert corrected["q2"].fill.start_color.rgb.endswith("E3F6EB")
    assert [c.value for c in book["Results 3"][1]] == header


def test_pdf_table_and_sheets(tmp_path):
    pytest.importorskip("reportlab")
    out, report = run(tmp_path, "pdf", {"pdf": {"mode": "table"}})
    assert out.read_bytes()[:4] == b"%PDF" and report["rows"] == 3
    image = np.full((200, 150), 255, np.uint8)
    out, report = run(
        tmp_path,
        "pdf",
        {"pdf": {"mode": "both"}},
        image_provider=lambda r: image,
        total=3,
    )
    assert out.read_bytes()[:4] == b"%PDF"
    with pytest.raises(ExportError):
        run(tmp_path, "pdf", {"pdf": {"mode": "sheets", "maxSheets": 2}}, total=3)
    # Many columns are split into bands rather than squeezed
    wide = {**INFO, "output_columns": [f"q{i}" for i in range(1, 120)]}
    out = tmp_path / "wide.pdf"
    export_results(
        lambda: iter(RESULTS), "pdf", out, ExportProfile.from_dict({}), [wide]
    )
    assert out.stat().st_size > 1000


def test_sqlite_upsert_and_schema_growth(tmp_path):
    profile = {
        "fields": [
            {"field": "Roll"},
            {"field": "price", "type": "decimal"},
            {"field": "dob", "type": "date", "format": "%d%m%Y"},
            {"field": "agree", "type": "bool"},
        ],
        "sql": {"table": "exam_results", "batchSize": 2},
    }
    out, report = run(tmp_path, "sqlite", profile)
    assert report["rows"] == 3
    out2, _ = run(tmp_path, "sqlite", profile)  # same file again: upsert, no duplicates
    assert out2 == out
    conn = sqlite3.connect(str(out))
    rows = conn.execute(
        "SELECT scan_id, job_id, Roll, price, dob, agree FROM exam_results ORDER BY scan_id"
    ).fetchall()
    assert len(rows) == 3
    assert rows[0] == ("scan0000", "job1", "012", 12.5, "2001-12-31", 1)
    types = {row[1]: row[2] for row in conn.execute("PRAGMA table_info(exam_results)")}
    assert (
        types["Roll"] == "TEXT"
        and types["price"] == "NUMERIC"
        and types["agree"] == "INTEGER"
    )
    conn.close()
    # A profile with a new column adds it to the existing table
    profile["fields"].append({"field": "q1"})
    run(tmp_path, "sqlite", profile)
    conn = sqlite3.connect(str(out))
    assert conn.execute("SELECT COUNT(*), MIN(q1) FROM exam_results").fetchone() == (
        3,
        "A",
    )
    conn.close()


def test_sqlalchemy_upsert(tmp_path):
    pytest.importorskip("sqlalchemy")
    target = tmp_path / "db.sqlite"
    profile = {
        "fields": [{"field": "Roll"}, {"field": "price", "type": "decimal"}],
        "sql": {"table": "results"},
    }
    for _ in range(2):
        out = tmp_path / "log.txt"
        report = export_results(
            lambda: iter(RESULTS),
            "sql",
            out,
            ExportProfile.from_dict(profile),
            [INFO],
            sql_url=f"sqlite+pysqlite:///{target}",
        )
        assert report["rows"] == 3
    conn = sqlite3.connect(str(target))
    assert conn.execute("SELECT COUNT(*) FROM results").fetchone()[0] == 3
    assert (
        conn.execute("SELECT price FROM results WHERE scan_id='scan0000'").fetchone()[0]
        == 12.5
    )
    conn.close()


def test_api_exports_job_with_corrections(tmp_path):
    spec = default_spec(questions=8, roll_digits=3, with_zones=False)
    template = spec.to_template(pre_processors=[])
    template["customLabels"] = {"RollNo": ["roll1..3"]}
    rng = random.Random(5)
    with make_client(tmp_path) as client:
        response = client.post(
            "/templates",
            files=[
                ("files", ("template.json", json.dumps(template), "application/json"))
            ],
            data={"name": "Exp"},
        )
        template_id = response.json()["id"]
        files = []
        for i in range(3):
            answers = random_answers(spec, rng, blank_rate=0.0)
            answers.update({"roll1": "0", "roll2": str(i), "roll3": "7"})
            image, _ = render_sheet(spec, answers, rng=rng)
            files.append(("files", (f"s{i}.png", png_bytes(image), "image/png")))
        job = client.post(
            "/jobs", data={"template_id": template_id}, files=files
        ).json()
        job = wait_for_job(client, job["id"])
        first = client.get(f"/results?job_id={job['id']}").json()["items"][0]["id"]
        client.post(f"/scans/{first}/corrections", json={"changes": {"q1": "D"}})

        body = {
            "format": "xlsx",
            "filters": {"job_id": job["id"]},
            "wait": True,
            "profile": {
                "fields": [{"field": "RollNo", "header": "Roll"}, {"field": "q1"}]
            },
        }
        record = client.post("/exports", json=body).json()
        assert record["state"] == "completed", record
        assert record["rows"] == 3
        download = client.get(record["download_url"])
        assert download.status_code == 200
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(download.content))
        rows = list(book.active.values)
        header = list(rows[0])
        assert header[-2:] == ["Roll", "q1"]
        assert [r[header.index("Roll")] for r in rows[1:]] == ["007", "017", "027"]
        assert rows[1][header.index("q1")] == "D"  # the corrected value
        assert rows[1][header.index("corrected")] is True

        # Lossy cast: refused, nothing to download
        body = {
            "format": "csv",
            "filters": {"job_id": job["id"]},
            "wait": True,
            "profile": {"fields": [{"field": "RollNo", "type": "int"}]},
        }
        record = client.post("/exports", json=body).json()
        assert record["state"] == "failed" and "leading zero" in record["error"]
        assert client.get(f"/exports/{record['id']}/download").status_code == 409

        preview = client.post(
            "/exports/preview",
            json={**body, "profile": {"fields": [{"field": "RollNo"}]}},
        ).json()
        assert preview["total"] == 3 and preview["rows"][0][-1] == "007"

        # Saved profiles; background export
        assert (
            client.put(
                "/export-profiles/My profile", json={"fields": [{"field": "q1"}]}
            ).status_code
            == 200
        )
        assert (
            client.get("/export-profiles").json()["profiles"][0]["name"] == "My profile"
        )
        record = client.post(
            "/exports",
            json={
                "format": "sqlite",
                "filters": {"job_id": job["id"]},
                "profile_name": "My profile",
            },
        ).json()
        for _ in range(100):
            state = client.get(f"/exports/{record['id']}").json()
            if state["state"] in ("completed", "failed"):
                break
            import time

            time.sleep(0.1)
        assert state["state"] == "completed", state
        assert len(client.get("/exports").json()["exports"]) == 3
        assert client.delete(f"/exports/{record['id']}").status_code == 200
        assert client.post("/exports", json={"format": "doc"}).status_code == 422


def test_cli_data_dir_and_jsonl(tmp_path):
    spec = default_spec(questions=4, roll_digits=2, with_zones=False)
    template_dir = tmp_path / "tpl"
    template_dir.mkdir()
    (template_dir / "template.json").write_text(
        json.dumps(spec.to_template(pre_processors=[]))
    )
    rng = random.Random(9)
    images = tmp_path / "in"
    images.mkdir()
    for i in range(2):
        image, _ = render_sheet(
            spec, random_answers(spec, rng, blank_rate=0.0), rng=rng
        )
        cv2.imwrite(str(images / f"s{i}.png"), image)
    from src.batch import main as batch_main

    out_dir = tmp_path / "batch"
    assert (
        batch_main(
            [
                "--template",
                str(template_dir / "template.json"),
                "--input",
                str(images),
                "--out",
                str(out_dir),
                "--workers",
                "1",
            ]
        )
        == 0
    )
    csv_out = tmp_path / "x.csv"
    assert (
        export_cli(
            [
                "--jsonl",
                str(out_dir / "results.jsonl"),
                "--template",
                str(template_dir / "template.json"),
                "--format",
                "csv",
                "--out",
                str(csv_out),
            ]
        )
        == 0
    )
    rows = list(csv.DictReader(io.StringIO(csv_out.read_text("utf-8-sig"))))
    assert len(rows) == 2 and "q1" in rows[0]
    pdf_out = tmp_path / "x.pdf"
    profile = tmp_path / "p.json"
    profile.write_text(json.dumps({"pdf": {"mode": "sheets"}}))
    assert (
        export_cli(
            [
                "--jsonl",
                str(out_dir / "results.jsonl"),
                "--template",
                str(template_dir / "template.json"),
                "--format",
                "pdf",
                "--profile",
                str(profile),
                "--out",
                str(pdf_out),
            ]
        )
        == 0
    )
    assert pdf_out.read_bytes()[:4] == b"%PDF"

    # --data-dir: export what the API stored
    with make_client(tmp_path) as client:
        response = client.post(
            "/templates",
            files=[
                (
                    "files",
                    (
                        "template.json",
                        (template_dir / "template.json").read_text(),
                        "application/json",
                    ),
                )
            ],
        )
        template_id = response.json()["id"]
        client.post(
            "/scans",
            data={"template_id": template_id},
            files=[("files", ("a.png", (images / "s0.png").read_bytes(), "image/png"))],
        )
    lossy = tmp_path / "lossy.json"
    lossy.write_text(
        json.dumps({"fields": [{"field": "q1", "type": "int"}], "strictCast": True})
    )
    xlsx_out = tmp_path / "x.xlsx"
    assert (
        export_cli(
            [
                "--data-dir",
                str(tmp_path / "data"),
                "--template-id",
                template_id,
                "--format",
                "xlsx",
                "--out",
                str(xlsx_out),
            ]
        )
        == 0
    )
    assert xlsx_out.exists()
    assert (
        export_cli(
            [
                "--data-dir",
                str(tmp_path / "data"),
                "--format",
                "csv",
                "--profile",
                str(lossy),
                "--out",
                str(tmp_path / "bad.csv"),
            ]
        )
        == 2
    )
    assert not (tmp_path / "bad.csv").exists()


def test_decimal_and_bool_options():
    from src.export.profile import Column, cast_value

    column = Column("p", "field", "p", type="decimal", options={"decimalComma": True})
    assert cast_value("1.234,50", column) == Decimal("1234.50")
    column = Column(
        "b", "field", "b", type="bool", options={"true": ["A"], "false": ["B"]}
    )
    assert cast_value("a", column) is True and cast_value("B", column) is False
