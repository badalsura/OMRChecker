"""
Streaming writers: CSV, XLSX (openpyxl write-only), PDF (reportlab) and SQL
(SQLite built in, other databases through SQLAlchemy). Each takes the column
list and an iterator of (values, states) rows and never holds all rows.
"""

import csv
import io
import re
import sqlite3
from decimal import Decimal

from src.export.profile import Column, ExportError, text_of
from src.export.rows import CORRECTED, FLAGGED, split_value

EXCEL_MAX_ROWS = 1048576  # including the header row
LEADING_ZERO_RE = re.compile(r"0\d+")
LONG_NUMBER_RE = re.compile(r"\d{16,}")  # Excel keeps 15 significant digits
CONTROL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _needs_text_guard(text):
    return bool(LEADING_ZERO_RE.fullmatch(text) or LONG_NUMBER_RE.fullmatch(text))


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def write_csv(path, columns, rows, options):
    mode = options.get("leadingZeros", "none")
    delimiter = options.get("delimiter", ",")
    encoding = "utf-8-sig" if options.get("bom", True) else "utf-8"
    count = 0
    with open(path, "w", newline="", encoding=encoding) as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        writer.writerow([c.key for c in columns])
        for values, _ in rows:
            out = []
            for column, value in zip(columns, values):
                text = text_of(value, column)
                if mode != "none" and column.type == "text" and _needs_text_guard(text):
                    # Keep Excel from turning 0123 into 123 (or 16+ digits into 1.2E+15)
                    text = f'="{text}"' if mode == "formula" else "'" + text
                out.append(text)
            writer.writerow(out)
            count += 1
    return count


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
def write_xlsx(path, columns, rows, options):
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise ExportError("XLSX export needs openpyxl (pip install openpyxl)") from None

    max_rows = int(options.get("maxRowsPerSheet") or EXCEL_MAX_ROWS - 1)
    max_rows = max(1, min(max_rows, EXCEL_MAX_ROWS - 1))
    base = str(options.get("sheetName") or "Results")[:25]
    highlight = options.get("highlight", True)
    fills = {
        FLAGGED: PatternFill("solid", start_color="FFF3C2", end_color="FFF3C2"),
        CORRECTED: PatternFill("solid", start_color="E3F6EB", end_color="E3F6EB"),
    }
    bold = Font(bold=True)
    formats = {}
    for index, column in enumerate(columns):
        if column.type == "date":
            formats[index] = column.options.get("excelFormat") or "yyyy-mm-dd"
        elif column.type == "text":
            formats[index] = "@"

    workbook = Workbook(write_only=True)
    sheet, sheet_rows, sheets, count = None, max_rows, 0, 0

    def new_sheet():
        nonlocal sheet, sheet_rows, sheets
        sheets += 1
        sheet = workbook.create_sheet(base if sheets == 1 else f"{base} {sheets}")
        sheet.freeze_panes = "A2"
        header = []
        for column in columns:
            cell = WriteOnlyCell(sheet, value=column.key)
            cell.font = bold
            header.append(cell)
        sheet.append(header)
        sheet_rows = 0

    for values, states in rows:
        if sheet_rows >= max_rows:
            new_sheet()
        out = []
        for index, (column, value, state) in enumerate(zip(columns, values, states)):
            if isinstance(value, str):
                value = CONTROL_CHARS_RE.sub("", value)
            styled = (highlight and state) or (
                isinstance(value, str) and value.startswith("=")
            )
            if not styled and column.type != "date":
                out.append(value)
                continue
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # never a formula
            if index in formats:
                cell.number_format = formats[index]
            if highlight and state:
                cell.fill = fills[state]
            out.append(cell)
        sheet.append(out)
        sheet_rows += 1
        count += 1
    if sheet is None:
        new_sheet()
    workbook.save(path)
    return count


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _reportlab():
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.pdfgen import canvas as pdf_canvas

        return A4, landscape, pdf_canvas
    except ImportError:
        raise ExportError(
            "PDF export needs reportlab (pip install reportlab)"
        ) from None


def _fit(c, text, width, font, size):
    if c.stringWidth(text, font, size) <= width:
        return text
    while text and c.stringWidth(text + "…", font, size) > width:
        text = text[:-1]
    return text + "…"


def pdf_table(c, columns, rows_factory, options, title):
    """Results table, split into column bands when it is too wide for a page."""
    A4, landscape, _ = _reportlab()
    page_w, page_h = landscape(A4)
    c.setPageSize((page_w, page_h))
    margin, size = 28, float(options.get("fontSize", 7))
    widths = [max(min(len(col.key), 18), 4) * size * 0.62 + 8 for col in columns]
    usable = page_w - 2 * margin
    bands, band, total = [], [], 0
    for index, width in enumerate(widths):
        if band and total + width > usable:
            bands.append(band)
            band, total = [], 0
        band.append(index)
        total += width
    if band:
        bands.append(band)
    row_h = size * 1.6
    count = 0
    for band_no, band in enumerate(bands):
        scale = usable / sum(widths[i] for i in band)
        band_widths = [widths[i] * min(scale, 2.5) for i in band]
        y = page_h

        def header():
            nonlocal y
            y = page_h - margin
            c.setFont("Helvetica-Bold", size + 2)
            suffix = f" (columns {band_no + 1}/{len(bands)})" if len(bands) > 1 else ""
            c.drawString(margin, y, title + suffix)
            y -= row_h * 1.6
            c.setFont("Helvetica-Bold", size)
            x = margin
            for i, w in zip(band, band_widths):
                c.drawString(
                    x + 2, y, _fit(c, columns[i].key, w - 4, "Helvetica-Bold", size)
                )
                x += w
            c.line(margin, y - 3, margin + sum(band_widths), y - 3)
            y -= row_h
            c.setFont("Helvetica", size)

        header()
        for values, states in rows_factory():
            if y < margin:
                c.showPage()
                header()
            x = margin
            for i, w in zip(band, band_widths):
                if states[i]:
                    c.setFillColorRGB(
                        *(
                            (1, 0.95, 0.76)
                            if states[i] == FLAGGED
                            else (0.89, 0.96, 0.92)
                        )
                    )
                    c.rect(x, y - 3, w, row_h, stroke=0, fill=1)
                    c.setFillColorRGB(0, 0, 0)
                c.drawString(
                    x + 2,
                    y,
                    _fit(c, text_of(values[i], columns[i]), w - 4, "Helvetica", size),
                )
                x += w
            y -= row_h
            if band_no == 0:
                count += 1
        c.showPage()
    return count


def pdf_sheets(c, columns, items, options, image_provider):
    """One page per sheet: the aligned image with the overlay, and the values."""
    A4, _, _ = _reportlab()
    from reportlab.lib.utils import ImageReader

    page_w, page_h = A4
    c.setPageSize((page_w, page_h))
    margin = 24
    count = 0
    for result, values, states in items:
        image = image_provider(result) if image_provider else None
        title = str(result.get("file_id") or result.get("scan_id"))
        c.setFont("Helvetica-Bold", 10)
        c.drawString(
            margin,
            page_h - margin,
            _fit(c, title, page_w - 2 * margin, "Helvetica-Bold", 10),
        )
        top = page_h - margin - 14
        image_w = (page_w - 2 * margin) * 0.62
        if image is not None:
            import cv2

            h, w = image.shape[:2]
            scale = min(image_w / w, (top - margin) / h)
            dw, dh = w * scale, h * scale
            ok, buffer = cv2.imencode(".jpg", image, [cv2.IMWRITE_JPEG_QUALITY, 80])
            if ok:
                c.drawImage(
                    ImageReader(io.BytesIO(buffer.tobytes())), margin, top - dh, dw, dh
                )
            _draw_overlay(c, result, margin, top, scale)
        else:
            c.setFont("Helvetica", 9)
            c.drawString(
                margin,
                top - 20,
                "No image available (source file missing and none stored)",
            )
        x = margin + image_w + 10
        width = page_w - margin - x
        y = top - 4
        size = 7.5
        for column, value, state in zip(columns, values, states):
            if y < margin:
                c.showPage()
                c.setPageSize((page_w, page_h))
                y = page_h - margin
                x, width = margin, page_w - 2 * margin
            if state:
                c.setFillColorRGB(
                    *((1, 0.95, 0.76) if state == FLAGGED else (0.89, 0.96, 0.92))
                )
                c.rect(x - 2, y - 2.5, width + 4, size * 1.45, stroke=0, fill=1)
                c.setFillColorRGB(0, 0, 0)
            c.setFont("Helvetica-Bold", size)
            c.drawString(
                x, y, _fit(c, column.key, width * 0.45, "Helvetica-Bold", size)
            )
            c.setFont("Helvetica", size)
            c.drawString(
                x + width * 0.47,
                y,
                _fit(c, text_of(value, column), width * 0.53, "Helvetica", size),
            )
            y -= size * 1.5
        c.showPage()
        count += 1
    return count


def _draw_overlay(c, result, left, top, scale):
    def rect(x, y, w, h, fill=False):
        c.rect(
            left + x * scale,
            top - (y + h) * scale,
            w * scale,
            h * scale,
            stroke=1,
            fill=1 if fill else 0,
        )

    pending = {item["name"] for item in result.get("review") or []}
    c.setLineWidth(0.4)
    for name, field in (result.get("fields") or {}).items():
        bubbles = field.get("bubbles") or []
        marks = set(
            split_value(field.get("value", ""), [b["value"] for b in bubbles]) or []
        )
        for b in bubbles:
            if b["value"] in marks:
                c.setStrokeColorRGB(0.18, 0.43, 0.87)
                c.setFillColorRGB(0.18, 0.43, 0.87, alpha=0.35)
                rect(b["x"], b["y"], b["w"], b["h"], fill=True)
            else:
                c.setStrokeColorRGB(
                    *((0.88, 0.54, 0) if name in pending else (0.6, 0.6, 0.6))
                )
                rect(b["x"], b["y"], b["w"], b["h"])
    for name, zone in (result.get("zones") or {}).items():
        box = zone.get("box") or []
        if len(box) == 4:
            c.setStrokeColorRGB(
                *((0.88, 0.54, 0) if name in pending else (0.12, 0.55, 0.12))
            )
            rect(*box)
    c.setFillColorRGB(0, 0, 0)
    c.setStrokeColorRGB(0, 0, 0)


# ---------------------------------------------------------------------------
# SQL
# ---------------------------------------------------------------------------
SQLITE_TYPES = {
    "text": "TEXT",
    "int": "INTEGER",
    "decimal": "NUMERIC",
    "date": "TEXT",
    "bool": "INTEGER",
}
IDENTIFIER_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_ .-]{0,62}$")


def sql_columns(columns):
    """The upsert key (scan_id) and job_id are always present."""
    columns = list(columns)
    names = {c.name for c in columns if c.source == "meta"}
    for key in ("job_id", "scan_id"):
        if key not in names:
            columns.insert(0, Column(key, "meta", key))
    seen = set()
    for column in columns:
        if not IDENTIFIER_RE.match(column.key):
            raise ExportError(
                f"Column name '{column.key}' is not usable in SQL; rename it in the profile"
            )
        if column.key.lower() in seen:
            raise ExportError(
                f"Column '{column.key}' appears twice (names are case-insensitive in SQL)"
            )
        seen.add(column.key.lower())
    return columns


def _key_index(columns):
    return next(
        i for i, c in enumerate(columns) if c.source == "meta" and c.name == "scan_id"
    )


def _sql_value(value):
    if isinstance(value, Decimal):
        return str(value)
    if hasattr(value, "isoformat") and not isinstance(value, str):
        return value.isoformat()
    if isinstance(value, bool):
        return int(value)
    return value


def _quote(name):
    return '"' + name.replace('"', '""') + '"'


def write_sqlite(path, table, columns, rows, batch_size=1000):
    if not IDENTIFIER_RE.match(table or ""):
        raise ExportError(f"Invalid table name '{table}'")
    key = columns[_key_index(columns)].key
    conn = sqlite3.connect(str(path))
    try:
        exists = conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)
        ).fetchone()
        if not exists:
            defs = [
                f"{_quote(c.key)} {SQLITE_TYPES[c.type]}"
                + (" PRIMARY KEY" if c.key == key else "")
                for c in columns
            ]
            conn.execute(f"CREATE TABLE {_quote(table)} ({', '.join(defs)})")
        else:
            present = {
                row[1].lower()
                for row in conn.execute(f"PRAGMA table_info({_quote(table)})")
            }
            for c in columns:
                if c.key.lower() not in present:
                    conn.execute(
                        f"ALTER TABLE {_quote(table)} ADD COLUMN {_quote(c.key)} {SQLITE_TYPES[c.type]}"
                    )
            conn.execute(
                f"CREATE UNIQUE INDEX IF NOT EXISTS {_quote('ux_' + table + '_' + key)} "
                f"ON {_quote(table)} ({_quote(key)})"
            )
        names = ", ".join(_quote(c.key) for c in columns)
        updates = ", ".join(
            f"{_quote(c.key)}=excluded.{_quote(c.key)}" for c in columns if c.key != key
        )
        sql = (
            f"INSERT INTO {_quote(table)} ({names}) VALUES ({', '.join('?' * len(columns))}) "
            f"ON CONFLICT({_quote(key)}) DO UPDATE SET {updates}"
        )
        count, batch = 0, []
        for values, _ in rows:
            batch.append([_sql_value(v) for v in values])
            if len(batch) >= batch_size:
                conn.executemany(sql, batch)
                conn.commit()
                count += len(batch)
                batch = []
        if batch:
            conn.executemany(sql, batch)
            count += len(batch)
        conn.commit()
        return count
    finally:
        conn.close()


def write_sqlalchemy(url, table, columns, rows, batch_size=1000):
    """Any SQLAlchemy URL (postgresql+psycopg://..., mysql+pymysql://...)."""
    try:
        import sqlalchemy as sa
    except ImportError:
        raise ExportError(
            "Exporting to this database needs SQLAlchemy and its driver "
            "(pip install SQLAlchemy psycopg[binary])"
        ) from None
    if not IDENTIFIER_RE.match(table or ""):
        raise ExportError(f"Invalid table name '{table}'")
    types = {
        "text": sa.Text,
        "int": sa.BigInteger,
        "decimal": sa.Numeric,
        "date": sa.Date,
        "bool": sa.Boolean,
    }
    key = columns[_key_index(columns)].key
    try:
        engine = sa.create_engine(url)
    except Exception as error:
        raise ExportError(f"Cannot use database URL: {error}") from None
    metadata = sa.MetaData()
    table_obj = sa.Table(
        table,
        metadata,
        *[
            sa.Column(
                c.key,
                sa.String(64) if c.key == key else types[c.type](),
                primary_key=c.key == key,
            )
            for c in columns
        ],
    )
    try:
        metadata.create_all(engine, checkfirst=True)
        present = {col["name"].lower() for col in sa.inspect(engine).get_columns(table)}
        preparer = engine.dialect.identifier_preparer
        with engine.begin() as conn:
            for c in columns:
                if c.key.lower() not in present:
                    kind = types[c.type]().compile(dialect=engine.dialect)
                    conn.execute(
                        sa.text(
                            f"ALTER TABLE {preparer.quote(table)} ADD COLUMN {preparer.quote(c.key)} {kind}"
                        )
                    )
        dialect = engine.dialect.name

        def flush(batch):
            if not batch:
                return
            with engine.begin() as conn:
                if dialect in ("postgresql", "sqlite"):
                    if dialect == "postgresql":
                        from sqlalchemy.dialects.postgresql import insert
                    else:
                        from sqlalchemy.dialects.sqlite import insert
                    statement = insert(table_obj)
                    statement = statement.on_conflict_do_update(
                        index_elements=[key],
                        set_={
                            c.key: statement.excluded[c.key]
                            for c in columns
                            if c.key != key
                        },
                    )
                    conn.execute(statement, batch)
                elif dialect in ("mysql", "mariadb"):
                    from sqlalchemy.dialects.mysql import insert

                    statement = insert(table_obj)
                    statement = statement.on_duplicate_key_update(
                        {
                            c.key: statement.inserted[c.key]
                            for c in columns
                            if c.key != key
                        }
                    )
                    conn.execute(statement, batch)
                else:
                    keys = [row[key] for row in batch]
                    conn.execute(table_obj.delete().where(table_obj.c[key].in_(keys)))
                    conn.execute(table_obj.insert(), batch)

        count, batch = 0, []
        for values, _ in rows:
            row = {}
            for c, v in zip(columns, values):
                if c.type == "decimal" and v is not None and not isinstance(v, Decimal):
                    v = Decimal(str(v))
                row[c.key] = v
            batch.append(row)
            if len(batch) >= batch_size:
                flush(batch)
                count += len(batch)
                batch = []
        flush(batch)
        count += len(batch)
        return count
    except ExportError:
        raise
    except Exception as error:
        raise ExportError(f"Database export failed: {error}") from None
    finally:
        engine.dispose()
